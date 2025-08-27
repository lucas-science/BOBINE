import os
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

from utils.pigna_constants import (
    TIME,
    TT301,
    TT302,
    TT303,
    FT240,
    PI177,
    PT230,
    DATA_REQUIRED,
    GRAPHS,
    TEMPERATURE_DEPENDING_TIME,
    DEBIMETRIC_RESPONSE_DEPENDING_TIME,
    PRESSURE_PYROLYSEUR_DEPENDING_TIME,
    PRESSURE_POMPE_DEPENDING_TIME,
    DELTA_PRESSURE_DEPENDING_TIME
)


class PignaData:
    def __init__(self, dir_root: str):
        self.first_file = ""
        if os.path.exists(dir_root):
            files = [f for f in os.listdir(dir_root)
                     if os.path.isfile(os.path.join(dir_root, f))
                     and not f.startswith('.')
                     and not f.startswith('~')
                     and not f.startswith('.~lock')
                     and f.lower().endswith('.csv')]

            if not files:
                raise FileNotFoundError(
                    f"Aucun fichier CSV valide trouvé dans {dir_root}")

            files.sort()
            self.first_file = os.path.join(dir_root, files[0])
        else:
            raise FileNotFoundError(f"Le répertoire {dir_root} n'existe pas")

        self.data_frame = pd.read_csv(self.first_file)
        self.columns = self.data_frame.columns.tolist()
        self.missing_columns = set(DATA_REQUIRED) - set(self.columns)

    def _select_columns(self, columns: list[str]) -> pd.DataFrame:
        return self.data_frame[columns]

    # --------------------------------------------------
    # Export data et services manquants
    # --------------------------------------------------

    def is_all_required_data(self) -> bool:
        return all(col in self.data_frame.columns for col in DATA_REQUIRED)

    def get_available_graphs(self) -> list[dict]:
        graphs = []
        for graph in GRAPHS:
            if all(col in self.columns for col in graph['columns']):
                graph['available'] = True
            else:
                graph['available'] = False
            graphs.append(graph)
        return graphs

    # --------------------------------------------------
    # Gestion des données manquantes
    # --------------------------------------------------

    def report_missing_per_column(self) -> pd.Series:
        return self.data_frame.isna().sum()

    def report_missing_per_row(self) -> pd.DataFrame:
        df = self.data_frame
        count_na = df.isna().sum(axis=1)
        cols_na = df.isna().apply(lambda row: list(row[row].index), axis=1)
        return pd.DataFrame({
            'n_missing': count_na,
            'cols_missing': cols_na
        }, index=df.index)

    # --------------------------------------------------
    # Extraction privé
    # --------------------------------------------------

    def _get_temperature_over_time(self) -> pd.DataFrame:
        cols = [TIME, TT301, TT302, TT303]
        return self._select_columns(cols).copy()

    def _get_debimetrique_response_over_time(self) -> pd.DataFrame:
        cols = [TIME, FT240]
        return self._select_columns(cols).copy()

    def _get_pression_pyrolyseur_over_time(self) -> pd.DataFrame:
        cols = [TIME, PI177]
        return self._select_columns(cols).copy()

    def _get_pression_sortie_pompe_over_time(self) -> pd.DataFrame:
        cols = [TIME, PT230]
        return self._select_columns(cols).copy()

    def _get_delta_pression_over_time(self) -> pd.DataFrame:
        cols = [TIME, PI177, PT230]
        df_sel = self._select_columns(cols).copy()
        delta_name = f"Delta_Pression_{PI177}_minus_{PT230}"
        df_sel[delta_name] = df_sel[PI177] - df_sel[PT230]
        return df_sel.drop(columns=[PI177, PT230])

    # --------------------------------------------------
    # Extraction publique
    # --------------------------------------------------

    def get_json_metrics(self, metric: str):
        if metric == TEMPERATURE_DEPENDING_TIME:
            return {
                "name": TEMPERATURE_DEPENDING_TIME,
                "data": self._get_temperature_over_time(),
                "x_axis": TIME,
                "y_axis": [TT301, TT302, TT303]
            }
        elif metric == DEBIMETRIC_RESPONSE_DEPENDING_TIME:
            return {
                "name": DEBIMETRIC_RESPONSE_DEPENDING_TIME,
                "data": self._get_debimetrique_response_over_time(),
                "x_axis": TIME,
                "y_axis": [FT240]
            }
        elif metric == PRESSURE_PYROLYSEUR_DEPENDING_TIME:
            return {
                "name": PRESSURE_PYROLYSEUR_DEPENDING_TIME,
                "data": self._get_pression_pyrolyseur_over_time(),
                "x_axis": TIME,
                "y_axis": [PI177]
            }
        elif metric == PRESSURE_POMPE_DEPENDING_TIME:
            return {
                "name": PRESSURE_POMPE_DEPENDING_TIME,
                "data": self._get_pression_sortie_pompe_over_time(),
                "x_axis": TIME,
                "y_axis": [PT230]
            }
        elif metric == DELTA_PRESSURE_DEPENDING_TIME:
            return {
                "name": DELTA_PRESSURE_DEPENDING_TIME,
                "data": self._get_delta_pression_over_time(),
                "x_axis": TIME,
                "y_axis": [f"Delta_Pression_{PI177}_minus_{PT230}"]
            }
        else:
            raise ValueError(f"Metric '{metric}' is not recognized.")

    def generate_workbook_with_charts(self,
        wb: Workbook,
        metrics_wanted: list[str],
        sheet_name="Pignat"
    ) -> Workbook:
        for metric in metrics_wanted:
            try:
                metric_data = self.get_json_metrics(metric)
                df = metric_data['data']
                # Création d'une feuille de données
                ws = wb.create_sheet(title=sheet_name)
                ws.append(df.columns.tolist())
                for row in df.itertuples(index=False):
                    ws.append(list(row))

                # Configuration du graphique
                chart = LineChart()
                chart.title = metric_data['name']
                chart.style = 13
                chart.y_axis.title = (
                    ', '.join(metric_data['y_axis'])
                    if len(metric_data['y_axis']) > 1
                    else metric_data['y_axis'][0]
                )
                chart.x_axis.title = metric_data['x_axis']

                max_row = ws.max_row
                # Séries de données (colonnes 2…n)
                data_ref = Reference(ws,
                                     min_col=2,
                                     min_row=1,
                                     max_col=1 + len(metric_data['y_axis']),
                                     max_row=max_row)
                chart.add_data(data_ref, titles_from_data=True)
                # Axe des abscisses (catégories)
                cats = Reference(ws, min_col=1, min_row=2, max_row=max_row)
                chart.set_categories(cats)

                # Insertion du graphique
                insert_col = chr(65 + df.shape[1] + 1)
                ws.add_chart(chart, f"{insert_col}2")

            except Exception as e:
                print(
                    f"Erreur traitement métrique '{metric}': {e}", file=sys.stderr)

        return wb
